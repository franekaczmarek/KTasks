"""Excel (openpyxl) and PDF (reportlab) summary reports built from the dashboard data."""
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NAVY, BERRY = "00205B", "D0006F"
SLA_LABEL = {"on_track": "On track", "at_risk": "At risk", "breached": "Breached", "met": "Met", "n_a": "N/A"}


def _fmt(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    return "" if v is None else v


def _scope_text(d: dict[str, Any]) -> str:
    s = d["scope"]
    period = f"last {s['days']} days" if s["days"] else "all time"
    vis = {"hidden": " · Hidden issues only", "visible": " · Visible issues only"}.get(s.get("visibility"), "")
    return f"Area: {s['area'] or 'All'} · Period: {period}{vis} · Issues in scope: {s['issue_count']}"


def kpi_rows(d: dict[str, Any]) -> list[tuple[str, Any]]:
    k, b = d["kpis"], d["blockers"]
    return [
        ("Open issues", k["open_issues"]),
        ("Open & blocked", k["open_blocked"]),
        ("Open & SLA breached", k["open_breached"]),
        ("Avg lead response (business days)", _fmt(k["avg_lead_response_days"])),
        ("Awaiting lead response", k["awaiting_response"]),
        ("SLA compliance % (resolved/closed)", _fmt(k["sla_compliance_pct"])),
        ("Resolved/closed issues", k["sla_finished_count"]),
        ("Total time lost to blockers (business days)", b["total_days"]),
        ("Active blockers", b["active"]),
    ]


def sla_status_rows(d: dict[str, Any]) -> list[tuple[str, str, int]]:
    active, finished = d["sla_status"]["active"], d["sla_status"]["finished"]
    return ([("Active (New / In Progress)", SLA_LABEL[k], v) for k, v in active.items()]
            + [("Resolved / Closed", SLA_LABEL[k], v) for k, v in finished.items()])


def issue_rows(d: dict[str, Any]) -> list[list[Any]]:
    return [[
        f"KT-{i['id']}", i["title"], i["area"], i["priority"], i["effort"], i["status"], i["lead_name"] or "",
        i["creator_name"], _fmt(i["created_at"]), i["metrics"]["lead_time_days"], i["metrics"]["blocker_time_days"],
        SLA_LABEL[i["metrics"]["sla_state"]], i["root_cause"] or "",
        "Yes" if i["metrics"]["is_blocked"] else "", _fmt(i["closed_at"]), i["closed_by_name"] or (
            "auto" if i["status"] == "Closed" else ""), i["rejected_reason"] or "", _fmt(i["expected_end_date"]) or "",
        "Yes" if i.get("is_hidden") else "",
    ] for i in d["issues"]]


ISSUE_HEADERS = ["ID", "Title", "Area", "Priority", "Effort", "Status", "Lead", "Reported by", "Created",
                 "Lead time (bd)", "Blocked (bd)", "SLA", "Root cause", "Blocked now", "Closed at", "Closed by", "Rejected reason",
                 "Agreed due date", "Hidden"]


def to_xlsx(d: dict[str, Any]) -> bytes:
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=NAVY)

    def sheet(ws, headers, rows, widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for idx, h in enumerate(headers, 1):
            longest = max([len(str(h))] + [len(str(r[idx - 1])) for r in rows] or [10])
            ws.column_dimensions[get_column_letter(idx)].width = (widths or {}).get(idx, min(max(longest + 2, 10), 60))

    ws = wb.active
    ws.title = "Summary"
    ws.append(["KTasks · Issue & Improvement Report"])
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws.append([_scope_text(d)])
    ws.append([f"Generated {_fmt(d['scope']['generated_at'])} UTC"])
    ws.append([])
    ws.append(["KPI", "Value"])
    for c in ws[5]:
        c.font, c.fill = head_font, head_fill
    for row in kpi_rows(d):
        ws.append(list(row))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 46, 14

    sheet(wb.create_sheet("SLA status"), ["Group", "SLA phase", "Issues"], sla_status_rows(d))
    sheet(wb.create_sheet("Issues"), ISSUE_HEADERS, issue_rows(d))
    sheet(wb.create_sheet("Root causes"), ["Root cause", "Issues"],
          [(r["root_cause"], r["count"]) for r in d["root_causes"]])
    sheet(wb.create_sheet("Quick wins"), ["Priority (impact)", "Effort", "Open issues", "Quick win"],
          [(c["priority"], c["effort"], c["count"], "Yes" if c["quick_win"] else "") for c in d["quick_wins"]])
    sheet(wb.create_sheet("Blockers"), ["Reason", "Occurrences", "Business days lost"],
          [(r["reason"], r["count"], r["days"]) for r in d["blockers"]["top_reasons"]])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(d: dict[str, Any]) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="KTasks report")
    styles = getSampleStyleSheet()
    h1 = styles["Title"].clone("h1", textColor=colors.HexColor(f"#{NAVY}"), alignment=0, fontSize=18)
    h2 = styles["Heading2"].clone("h2", textColor=colors.HexColor(f"#{NAVY}"), spaceBefore=10)
    small = styles["Normal"].clone("small", fontSize=8, leading=10)
    cell = styles["Normal"].clone("cell", fontSize=7.5, leading=9)

    def table(rows, widths=None, header=True):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        style = [("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E5EB")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                                                        [colors.white, colors.HexColor("#F4F5F7")])]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                      ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
        t.setStyle(TableStyle(style))
        return t

    story = [
        Paragraph("KTasks · Issue &amp; Improvement Report", h1),
        Paragraph(f"{_scope_text(d)} · generated {_fmt(d['scope']['generated_at'])} UTC", small),
        Spacer(1, 6),
        Paragraph("Key indicators", h2),
        table([["KPI", "Value"]] + [[k, str(_fmt(v))] for k, v in kpi_rows(d)], [110 * mm, 30 * mm]),
        Paragraph(f"SLA status ({d['sla_status']['agreed_count']} active issues on an agreed due date, "
                  f"{d['sla_status']['pending_requests']} due date changes awaiting the reporter)", h2),
        table([["Group", "SLA phase", "Issues"]] + [list(r) for r in sla_status_rows(d)],
              [60 * mm, 30 * mm, 25 * mm]),
        Paragraph("Root cause analysis (resolved / closed)", h2),
        table([["Root cause", "Issues"]] + [[r["root_cause"], r["count"]] for r in d["root_causes"]],
              [60 * mm, 25 * mm]),
        Paragraph("Quick wins matrix: open issues by priority (impact) × effort", h2),
        table([["Priority \\ Effort"] + ["Low", "Medium", "High"]] + [
            [p] + [next(c["count"] for c in d["quick_wins"] if c["priority"] == p and c["effort"] == e)
                   for e in ("Low", "Medium", "High")]
            for p in ("Critical", "High", "Medium", "Low")], [40 * mm, 25 * mm, 25 * mm, 25 * mm]),
        Paragraph(f"Blockers: {d['blockers']['total_days']} business days lost, "
                  f"{d['blockers']['active']} active", h2),
        table([["Reason", "Occurrences", "Business days lost"]] +
              [[Paragraph(r["reason"], cell), r["count"], r["days"]] for r in d["blockers"]["top_reasons"]]
              or [["No blockers recorded", "", ""]], [150 * mm, 30 * mm, 35 * mm]),
        Paragraph("Issues", h2),
    ]
    cols = [0, 1, 2, 3, 5, 6, 9, 10, 11, 12]
    widths = [16, 88, 24, 18, 20, 30, 18, 18, 18, 24]
    rows = [[ISSUE_HEADERS[c] for c in cols]] + [
        [Paragraph(str(r[c]), cell) if c == 1 else r[c] for c in cols] for r in issue_rows(d)]
    story.append(table(rows, [w * mm for w in widths]))
    doc.build(story)
    return buf.getvalue()
