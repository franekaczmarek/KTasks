"""Hidden issues: invisible to uninvolved employees everywhere; staff filter all/hidden/visible."""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import DIRECTOR, EMPLOYEE, EMPLOYEE2, EMPLOYEE3, LEAD, LEAD2, auth
from tests.test_backup import ola_covers_marek
from tests.test_issues import actions


def ids(rows, key="id"):
    return {r[key] for r in rows}


def export_ids(client, who):
    r = client.get("/reports/export", params={"format": "xlsx"}, headers=auth(who))
    return {row[0] for row in load_workbook(BytesIO(r.content))["Issues"].iter_rows(min_row=2, values_only=True)}


def test_hidden_issue_is_invisible_to_uninvolved_employees(client, make_issue, db):
    issue = make_issue(title="Confidential payroll leak", summary="sensitive")
    iid = issue["id"]
    client.post(f"/issues/{iid}/join", headers=auth(EMPLOYEE2))                     # Piotr follows it first
    client.post(f"/issues/{iid}/tasks", json={"title": "Quarantine payroll export"}, headers=auth(LEAD))
    client.post(f"/issues/{iid}/comments", json={"content": "payroll comment"}, headers=auth(EMPLOYEE))

    r = client.patch(f"/issues/{iid}", json={"is_hidden": True}, headers=auth(LEAD))
    assert r.status_code == 200 and r.json()["is_hidden"] is True
    client.post(f"/issues/{iid}/comments", json={"content": "after hiding"}, headers=auth(LEAD))

    piotr = auth(EMPLOYEE2)
    assert client.get(f"/issues/{iid}", headers=piotr).status_code == 404
    assert iid not in ids(client.get("/issues", headers=piotr).json())
    assert iid not in ids(client.get("/issues", params={"visibility": "hidden"}, headers=piotr).json())
    found = client.get("/search", params={"q": "payroll"}, headers=piotr).json()
    assert iid not in ids(found["issues"]) and iid not in ids(found["tasks"], "issue_id")
    assert iid not in ids(found["discussions"], "issue_id")
    assert iid not in ids(client.get("/discussions", headers=piotr).json())
    assert iid not in ids(client.get("/tasks", headers=piotr).json(), "issue_id")
    assert iid not in ids(client.get("/issues/duplicates", params={"q": "Confidential payroll leak"},
                                     headers=piotr).json())
    assert iid not in ids(client.get("/notifications", params={"limit": 100}, headers=piotr).json()["items"],
                          "issue_id")
    for method, path, body in (("get", f"/issues/{iid}/comments", None), ("get", f"/issues/{iid}/activity", None),
                               ("post", f"/issues/{iid}/comments", {"content": "hi"}),
                               ("post", f"/issues/{iid}/blockers", {"reason": "nope"}),
                               ("post", f"/issues/{iid}/join", None)):
        assert client.request(method.upper(), path, json=body, headers=piotr).status_code == 404, path
    assert f"KT-{iid}" not in export_ids(client, EMPLOYEE2)
    # The hidden comment produced no notification row for Piotr at all.
    sent = db.execute("select count(*) from public.notifications n join public.users u on u.id = n.user_id "
                      "where u.email = %s and n.issue_id = %s and n.message like %s",
                      (EMPLOYEE2, iid, "%after hiding%")).fetchone()[0]
    assert sent == 0

    # The reporter and staff still see it, flagged.
    for who in (EMPLOYEE, LEAD2, DIRECTOR):
        r = client.get(f"/issues/{iid}", headers=auth(who))
        assert r.status_code == 200 and r.json()["is_hidden"] is True
    assert f"KT-{iid}" in export_ids(client, LEAD)
    assert actions(client, iid)[-1] == "issue_hidden"


def test_task_assignee_sees_hidden_issue(client, users, make_issue):
    iid = make_issue(email=LEAD, title="Lead-only hidden", hidden="true")["id"]
    assert client.get(f"/issues/{iid}", headers=auth(EMPLOYEE2)).status_code == 404
    client.post(f"/issues/{iid}/tasks", json={"title": "Piotr helps", "assignee_id": users[EMPLOYEE2]["id"]},
                headers=auth(LEAD))
    r = client.get(f"/issues/{iid}", headers=auth(EMPLOYEE2))
    assert r.status_code == 200 and r.json()["is_hidden"] is True
    assert iid in ids(client.get("/tasks", headers=auth(EMPLOYEE2)).json(), "issue_id")


def test_staff_visibility_filter(client, make_issue):
    hidden = make_issue(email=LEAD, title="Filter hidden one", hidden="true")["id"]
    shown = make_issue(title="Filter visible one")["id"]
    every = ids(client.get("/issues", headers=auth(LEAD)).json())
    only_hidden = client.get("/issues", params={"visibility": "hidden"}, headers=auth(LEAD)).json()
    only_visible = client.get("/issues", params={"visibility": "visible"}, headers=auth(LEAD)).json()
    assert {hidden, shown} <= every
    assert hidden in ids(only_hidden) and all(i["is_hidden"] for i in only_hidden)
    assert shown in ids(only_visible) and not any(i["is_hidden"] for i in only_visible)
    d = client.get("/dashboard", params={"visibility": "hidden"}, headers=auth(LEAD)).json()
    assert d["scope"]["visibility"] == "hidden" and d["scope"]["issue_count"] == len(only_hidden)
    assert hidden in ids(client.get("/discussions", params={"visibility": "hidden"}, headers=auth(LEAD)).json())
    assert shown not in ids(client.get("/discussions", params={"visibility": "hidden"}, headers=auth(LEAD)).json())


def test_only_staff_or_acting_lead_can_hide(client, make_issue):
    r = client.post("/issues", data={"title": "[pytest] Sneaky hide", "area": "Operations", "priority": "Low",
                                     "effort": "Low", "hidden": "true"}, headers=auth(EMPLOYEE))
    assert r.status_code == 403
    iid = make_issue(title="Reporter cannot hide")["id"]
    assert client.patch(f"/issues/{iid}", json={"is_hidden": True}, headers=auth(EMPLOYEE)).status_code == 403
    assert client.patch(f"/issues/{iid}", json={"visible_to_backup": True},
                        headers=auth(LEAD)).status_code == 422             # only hidden issues


def test_employee_backup_sees_hidden_issue_only_when_opened_to_them(client, users, make_issue):
    iid = make_issue(email=LEAD, title="Hidden in Improvements", area="Improvements", hidden="true")["id"]
    with ola_covers_marek(client, users, absent=False):
        assert client.get(f"/issues/{iid}", headers=auth(EMPLOYEE3)).status_code == 404
        r = client.patch(f"/issues/{iid}", json={"visible_to_backup": True}, headers=auth(LEAD2))
        assert r.status_code == 200 and r.json()["visible_to_backup"] is True
        assert client.get(f"/issues/{iid}", headers=auth(EMPLOYEE3)).status_code == 200
        assert client.get(f"/issues/{iid}", headers=auth(EMPLOYEE2)).status_code == 404   # still hidden for others

        # Unhiding makes it public and closes the special backup access.
        r = client.patch(f"/issues/{iid}", json={"is_hidden": False}, headers=auth(LEAD))
        assert r.json()["is_hidden"] is False and r.json()["visible_to_backup"] is False
        assert client.get(f"/issues/{iid}", headers=auth(EMPLOYEE2)).status_code == 200
    assert actions(client, iid)[-3:] == ["backup_access_granted", "issue_unhidden", "backup_access_revoked"]
